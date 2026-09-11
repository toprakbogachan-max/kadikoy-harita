#!/usr/bin/env python3
"""Mekan aday havuzu -> biçimlendirilmiş .xlsx

Kullanım:
    python3 havuz.py adaylar.json cikti.xlsx

adaylar.json şeması — bir liste, her eleman:

    {
      "mekan":   "Rafine Moda",                 # zorunlu
      "kategori":"Restoran",                    # Restoran / Kafe / Meyhane / Bar / Pastane ...
      "mutfak":  "Kahvaltı / fine cuisine",
      "semt":    "Moda",
      "adres":   "Caferağa, Nene Hatun Sk. No:3A",
      "puan":    4.5,                           # Google Haritalar
      "yorum":   1598,                          # Google Haritalar
      "kaynaklar": ["TikTok", "Gurman Atlas", "Oggusto"],   # kaynak sayısı buradan sayılır
      "fiyat":   "₺600–800",                    # boş olabilir
      "acik":    "Açık",                        # Açık | Geçici kapalı | Kalıcı kapalı
      "durum":   "Sırada",                      # Sırada | Araştırılacak | Elendi
      "not":     "Üç kaynakta birden geçiyor"
    }

Popülerlik, sosyal skor ve öncelik skoru FORMÜL olarak yazılır — hardcode edilmez.
Yazdıktan sonra recalc et (LibreOffice gerekir):
    python3 /mnt/skills/public/xlsx/scripts/recalc.py cikti.xlsx
    # yerelde:  soffice --headless --convert-to xlsx --outdir . cikti.xlsx
"""
import json
import sys
from datetime import date

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FONT = "Arial"
HEADERS = ["Mekan", "Kategori", "Mutfak / Alt tür", "Semt", "Adres",
           "Google puanı", "Yorum sayısı", "Kaynak sayısı", "Kaynaklar",
           "Popülerlik", "Sosyal skor", "Öncelik skoru",
           "Fiyat (kişi başı)", "Açık mı", "Durum", "Not"]
WIDTHS = [30, 13, 24, 22, 36, 12, 12, 12, 32, 13, 12, 13, 16, 13, 15, 56]
STATUS_FILL = {
    "Sırada": PatternFill("solid", fgColor="DDEBD8"),
    "Araştırılacak": PatternFill("solid", fgColor="FDF2D0"),
    "Elendi": PatternFill("solid", fgColor="F4DAD6"),
}


def build(items, out_path, tarih=None):
    tarih = tarih or date.today().strftime("%d.%m.%Y")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Aday Havuzu"

    hdr_fill = PatternFill("solid", fgColor="1D1B19")
    hdr_font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
    cell_font = Font(name=FONT, size=10)
    cell_bold = Font(name=FONT, size=10, bold=True)
    input_font = Font(name=FONT, size=10, color="0000FF")
    thin = Side(style="thin", color="D9D2C8")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.append(HEADERS)
    for it in items:
        ks = it.get("kaynaklar") or []
        ws.append([
            it.get("mekan", ""), it.get("kategori", ""), it.get("mutfak", ""),
            it.get("semt", ""), it.get("adres", ""),
            it.get("puan"), it.get("yorum"), len(ks), ", ".join(ks),
            "", "", "",                       # J, K, L formülle dolacak
            it.get("fiyat", ""), it.get("acik", ""), it.get("durum", ""), it.get("not", ""),
        ])

    first, last = 2, len(items) + 1
    for i in range(first, last + 1):
        ws["J%d" % i] = ('=IF(G{0}>=3000,"Çok yüksek",IF(G{0}>=1500,"Yüksek",'
                         'IF(G{0}>=700,"Orta","Düşük")))').format(i)
        ws["K%d" % i] = "=ROUND(MIN(H{0},3)/3*100,1)".format(i)
        ws["L%d" % i] = ('=IF(N{0}<>"Açık",0,ROUND((F{0}/5)*45'
                         '+(MIN(G{0},5000)/5000)*30+(MIN(H{0},3)/3)*25,1))').format(i)

    for c in range(1, len(HEADERS) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font, cell.fill, cell.border = hdr_font, hdr_fill, border
        cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
    ws.row_dimensions[1].height = 34

    for i in range(first, last + 1):
        for c in range(1, len(HEADERS) + 1):
            cell = ws.cell(row=i, column=c)
            cell.font, cell.border = cell_font, border
            cell.alignment = Alignment(vertical="center", wrap_text=(c in (5, 9, 16)))
        ws.cell(row=i, column=1).font = cell_bold
        ws.cell(row=i, column=6).number_format = "0.0"
        ws.cell(row=i, column=7).number_format = "#,##0"
        ws.cell(row=i, column=11).number_format = "0.0"
        ws.cell(row=i, column=12).number_format = "0.0"
        for c in (6, 7, 8):                    # mavi = elle girilen ham veri
            ws.cell(row=i, column=c).font = input_font
        st = ws.cell(row=i, column=15).value
        if st in STATUS_FILL:
            ws.cell(row=i, column=15).fill = STATUS_FILL[st]
        ws.row_dimensions[i].height = 32

    for c, w in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = "A1:P%d" % last

    note = last + 2
    ws.cell(row=note, column=1, value="TABLO HAKKINDA").font = Font(name=FONT, size=10, bold=True)
    lines = [
        "Kapsam: Moda ve Kadıköy. Kadıköy Harita'da hâlihazırda pinli olan mekanlar bu havuza alınmadı.",
        "Keşif kaynakları, öncelik sırasıyla: (1) TikTok hashtag'leri, (2) Instagram #kadikoyyemek, "
        "(3) Oggusto Moda listesi, (4) Gurman Atlas (Vedat Milor) Kadıköy listesi.",
        "Puan, yorum sayısı ve kaynak sayısı (mavi hücreler) elle girilmiş ham veridir. "
        "Puan ve yorum sayısı Google Haritalar, %s itibarıyla." % tarih,
        "Popülerlik: yorum sayısına göre otomatik bant — 3.000+ Çok yüksek · 1.500+ Yüksek · 700+ Orta · altı Düşük.",
        "Sosyal skor: mekanın kaç farklı kaynakta göründüğü, 100 üzerinden (3 ve üzeri kaynak = 100).",
        "Öncelik skoru: 100 üzerinden = (puan/5)×45 + (yorum sayısı/5.000, en fazla 1)×30 "
        "+ (kaynak sayısı/3, en fazla 1)×25. Kapalı mekanlar otomatik 0 alır.",
        "Durum: Sırada = fotoğraf aşamasına hazır · Araştırılacak = önce yorum analizi gerekiyor "
        "· Elendi = kapalı ya da kritere uymuyor.",
        "Yeni mekan eklerken A–I ve M–P sütunlarını doldur; J, K ve L kendiliğinden hesaplanır.",
    ]
    for k, t in enumerate(lines):
        ws.cell(row=note + 1 + k, column=1, value=t).font = Font(name=FONT, size=9, color="5B534B")

    wb.save(out_path)
    return len(items)


def main():
    if len(sys.argv) < 3:
        sys.exit(__doc__)
    with open(sys.argv[1], encoding="utf-8") as fh:
        items = json.load(fh)
    n = build(items, sys.argv[2])
    print("ok %d mekan -> %s" % (n, sys.argv[2]))


if __name__ == "__main__":
    main()
